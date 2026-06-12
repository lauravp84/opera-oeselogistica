# -*- coding: utf-8 -*-
import re, copy
from docx import Document

SRC = "Ementario_Operacoes_e_Apoio_a_Decisao_v2.docx"
DST = "Ementario_Operacoes_e_Apoio_a_Decisao_v3.docx"

doc = Document(SRC)

def all_paragraphs(d):
    for p in d.paragraphs:
        yield p
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    yield p

pat = re.compile(r"\*([^*]+)\*")

orig_text_no_ast = []
heading_count = 0

def fix_paragraph(p):
    full = "".join(r.text for r in p.runs)
    if "*" not in full:
        return
    # process run by run; asterisk pairs may span runs — handle per-run first
    # Build new run list by splitting runs on asterisk pairs across the paragraph.
    # Simple robust approach: rebuild runs, splitting each run's text by the pattern,
    # tracking italic state across runs (asterisks may span runs).
    italic_open = False
    new_runs = []  # (text, italic_flag, source_run)
    for r in p.runs:
        parts = r.text.split("*")
        for i, part in enumerate(parts):
            if i > 0:
                italic_open = not italic_open
            if part:
                new_runs.append((part, italic_open, r))
    # replace runs
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    for text, ital, src in new_runs:
        new_el = copy.deepcopy(src._element)
        p._p.append(new_el)
        from docx.text.run import Run
        nr = Run(new_el, p)
        nr.text = text
        if ital:
            nr.italic = True

paras = list(all_paragraphs(doc))
for p in paras:
    txt = "".join(r.text for r in p.runs)
    orig_text_no_ast.append(txt.replace("*", ""))
    if (p.style.name if p.style else "").startswith("Heading"):
        pass

# count discipline headings before (heading paragraphs matching ACA codes or heading style level)
orig_headings = [p.text for p in paras if (p.style.name if p.style else "").startswith("Heading")]

for p in paras:
    fix_paragraph(p)

doc.save(DST)

# validation
doc2 = Document(DST)
paras2 = list(all_paragraphs(doc2))
texts2 = ["".join(r.text for r in p.runs) for p in paras2]
remaining = sum(t.count("*") for t in texts2)
headings2 = [p.text for p in paras2 if (p.style.name if p.style else "").startswith("Heading")]
identical = texts2 == orig_text_no_ast

# discipline heading count: headings containing ACA code
disc = [h for h in headings2 if re.search(r"ACA\d{3}", h)]
disc_orig = [h for h in orig_headings if re.search(r"ACA\d{3}", h)]

print("asteriscos restantes:", remaining)
print("headings totais antes/depois:", len(orig_headings), len(headings2))
print("headings de disciplina (ACA###) antes/depois:", len(disc_orig), len(disc))
print("texto identico (sem asteriscos):", identical)
if not identical:
    for a, b in zip(orig_text_no_ast, texts2):
        if a != b:
            print("DIFF:\n <", repr(a), "\n >", repr(b))
            break
